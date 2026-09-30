import os
"""Export the app as an old-style workbook, import it unchanged (nothing to do),
then edit it like a person would in Excel and import that (exactly the edits)."""
from testbase import BASE, IMPORT_SNAP, KEEP
import io, sys, time
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import openpyxl
import pandas as pd
from sqlalchemy import text
from itemmaster.db import get_engine
from itemmaster import dept_mapping as dm, old_workbook_import as owi
E = get_engine(); F = []
SAFE = IMPORT_SNAP
SP = os.path.dirname(os.path.abspath(__file__))


def check(ok, msg):
    print(("  PASS " if ok else "  FAIL ") + msg)
    if not ok: F.append(msg)


def plan_of(data: bytes):
    ex = owi.extract(owi.read_workbook(io.BytesIO(data)))
    return owi.plan(E, ex, dm.get_departments(E).iloc[:, 0].tolist(), "AJ")


def todo(p):
    return {"groups": p["groups"][p["groups"]["Result"] == owi.STAGE], "items": p["items"][p["items"]["Result"] == owi.STAGE],
            "moves": p["moves"][p["moves"]["Result"] == owi.MOVE]}


try:
    dm.restore_snapshot(E, SAFE, "AJ")  # the old-workbook import staged
    t = time.time(); data = owi.export_workbook(E); print(f"  exported {len(data)/1e6:.1f} MB in {time.time()-t:.0f}s")
    open(SP + r"\Item Master department workbook (export).xlsx", "wb").write(data)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    counts = {ws.title: ws.max_row - 2 for ws in wb.worksheets if ws.title not in ("Start Here", "Departments")}
    print("  rows per sheet:", counts)
    with E.connect() as c:
        st = dict(c.execute(text("SELECT decision_state, COUNT(*) FROM dbo.dept_mapping_combos GROUP BY decision_state")).all())
    check(counts["Department Mapping Broken Out"] == st.get("broken_out", 0) and counts["Decided Broken Out Combos"] == st.get("decided_broken_out", 0),
          "every Broken Out and finished Broken Out group is on its sheet")
    check(sum(counts[s] for s in ("Department Mapping Crosswalk", "Department Mapping Unmatched", "Department Mapping Broken Out",
                                  "Department Mapping Decided", "Decided Broken Out Combos")) == sum(st.values()),
          f"every one of the app's {sum(st.values()):,} groups is on exactly one sheet")

    print("\n== 1. Import the export unchanged")
    t = time.time(); p = plan_of(data); print(f"  planned in {time.time()-t:.0f}s")
    d = todo(p)
    check(len(d["groups"]) == 0 and len(d["items"]) == 0 and len(d["moves"]) == 0,
          f"nothing to stage or move ({len(d['groups'])} groups, {len(d['items'])} items, {len(d['moves'])} moves)")
    done = p["groups"]["Result"].str.startswith(owi.DONE).sum()
    check(done == 22, f"the import's 22 staged group decisions come back as Already done ({done})")

    print("\n== 2. Edit it in Excel and import")
    busy = set(dm.get_pending_changes(E)) | {m["combo_id"] for m in dm.get_recent_moves(E)} | \
        {c["combo_id"] for c in dm.get_pending_upc_changes(E).values()}
    facts = pd.read_sql(text("SELECT combo_id, source_key, raw_department, raw_category, raw_subcategory, decision_state, "
                             "decided_department, n_upcs_total FROM dbo.dept_mapping_combos"), E)
    src = dict(pd.read_sql(text("SELECT source_key, source_label FROM dbo.sources"), E).itertuples(index=False))
    key = lambda r: (str(src.get(r.source_key) or r.source_key).upper(), *(owi._s(v).upper() for v in (r.raw_department, r.raw_category, r.raw_subcategory)))
    by_key = {key(r): r for r in facts.itertuples()}

    def sheet_rows(name):
        ws = wb[name]
        hdr = [c.value for c in ws[2]]
        return ws, hdr, [(i, {h: v for h, v in zip(hdr, [c.value for c in row])}) for i, row in enumerate(ws.iter_rows(min_row=3), start=3)]

    def pick(name, want_state, n_max=30):
        ws, hdr, rows = sheet_rows(name)
        for i, r in rows:
            k = tuple(owi._s(r.get(c)).upper() for c in ("Source", "Old Department", "Category", "Subcategory"))
            f = by_key.get(k)
            if f is not None and f.combo_id not in busy and f.decision_state == want_state and 2 <= f.n_upcs_total <= n_max:
                busy.add(f.combo_id)
                return ws, hdr, i, r, f
        raise RuntimeError("nothing to pick on " + name)

    def setc(ws, hdr, row, col, val):
        ws.cell(row=row, column=hdr.index(col) + 1, value=val)

    expect = {}
    ws, hdr, i, r, f = pick("Department Mapping Crosswalk", "not_reviewed")
    newd = "GROCERY" if r.get("New Department") != "GROCERY" else "DELI"
    setc(ws, hdr, i, "Action", "Approve"); setc(ws, hdr, i, "Manual Override Department", newd)
    expect[f.combo_id] = ("group", newd); print("   approve Crosswalk group", f.combo_id, "→", newd)
    ws, hdr, i, r, f = pick("Department Mapping Decided", "not_reviewed")
    newd = "DELI" if f.decided_department != "DELI" else "GROCERY"
    setc(ws, hdr, i, "Manual Override Department", newd)
    expect[f.combo_id] = ("group", newd); print("   change Decided group", f.combo_id, f.decided_department, "→", newd)
    ws, hdr, i, r, f = pick("Department Mapping Decided", "not_reviewed")
    setc(ws, hdr, i, "Action", "Send to Broken Out")
    expect[f.combo_id] = ("move", "broken_out"); print("   send Decided group", f.combo_id, "to Broken Out")
    ws, hdr, i, r, f = pick("Department Mapping Broken Out", "broken_out", 500)
    setc(ws, hdr, i, "Action", "Return to Crosswalk/Unmatched")
    expect[f.combo_id] = ("move", "not_reviewed"); print("   return Broken Out group", f.combo_id, "for review")
    ws, hdr, rows = sheet_rows("Department UPC Overrides")
    bo_pick = None
    for i, r in rows:
        k = tuple(owi._s(r.get(c)).upper() for c in ("Source", "Old Department", "Category", "Subcategory"))
        f = by_key.get(k)
        if f is not None and f.combo_id not in busy and r.get("Decided Via") == "Needs Review":
            bo_pick = bo_pick or f.combo_id
            if f.combo_id == bo_pick:
                setc(ws, hdr, i, "Manual Override Department", "BAKERY")
                expect.setdefault("items", []).append(r["UPC"])
                if len(expect["items"]) == 2:
                    break
    print("   decide 2 Broken Out items in group", bo_pick, expect["items"])
    with E.connect() as c:
        free_upc = c.execute(text("SELECT TOP 1 i.upc FROM dbo.items i WHERE i.department = 'GROCERY' AND i.upc NOT IN "
                                  "(SELECT upc FROM dbo.dept_mapping_combo_upcs) AND i.upc NOT IN (SELECT upc FROM dbo.manual_overrides)")).scalar()
    ws = wb["Final UPC Overrides"]
    ws.cell(row=ws.max_row + 1, column=1, value=free_upc); ws.cell(row=ws.max_row, column=2, value="DELI")
    print("   add a UPC override", free_upc, "→ DELI")
    buf = io.BytesIO(); wb.save(buf); edited = buf.getvalue()

    p = plan_of(edited); d = todo(p)
    staged_groups = dict(zip(d["groups"]["combo_id"], d["groups"]["Will be"]))
    check(staged_groups == {k: v[1] for k, v in expect.items() if isinstance(k, int) and v[0] == "group"},
          f"exactly the 2 group edits are staged {staged_groups}")
    moved = dict(zip(d["moves"]["combo_id"], d["moves"]["move"]))
    check(moved == {k: {"broken_out": "to_broken_out", "not_reviewed": "to_review"}[v[1]] for k, v in expect.items() if isinstance(k, int) and v[0] == "move"},
          f"exactly the 2 moves {moved}")
    it = d["items"]
    check(sorted(it.loc[it["Goes to"].str.startswith("Broken Out"), "UPC"]) == sorted(expect["items"]), "the 2 Broken Out item decisions")
    check(list(it.loc[it["Goes to"] == "UPC override", "UPC"]) == [free_upc], "the 1 UPC override (from Final UPC Overrides)")
    before = dm.old_workbook_import_summary(E)
    r = owi.apply(E, p, "AJ", True, "edited export.xlsx")
    print("   applied:", r)
    with E.connect() as c:
        states = dict(c.execute(text("SELECT combo_id, decision_state FROM dbo.dept_mapping_combos WHERE combo_id IN :ids")
                                .bindparams(__import__("sqlalchemy").bindparam("ids", expanding=True)),
                                {"ids": [k for k in expect if isinstance(k, int)]}).all())
    check(all(states[k] == v[1] for k, v in expect.items() if isinstance(k, int) and v[0] == "move"), f"moves made {states}")
    pc = dm.get_pending_changes(E)
    check(all(pc.get(k, {}).get("department") == v[1] for k, v in expect.items() if isinstance(k, int) and v[0] == "group"), "group edits staged")
    pu = dm.get_pending_upc_changes(E); im = dm.get_item_master_pending(E)
    check(all(pu.get(u, {}).get("department") == "BAKERY" for u in expect["items"]) and im.get(free_upc, {}).get("department") == "DELI",
          "item decisions and the UPC override staged")
    after = dm.old_workbook_import_summary(E)
    check(after["groups"] == before["groups"] + 2 and after["items"] == before["items"] + 2, f"tagged as coming from a workbook upload {after}")
finally:
    print("\n  restoring #%d…" % SAFE); dm.restore_snapshot(E, SAFE, "AJ")
    for s_ in dm.list_snapshots(E).to_dict("records"):
        if str(s_["label"]).startswith("Before importing edited export"):
            dm.delete_snapshot(E, s_["snapshot_id"])
    print("  vs #%d:" % SAFE, dm.compare_snapshot_to_live(E, SAFE), dm.old_workbook_import_summary(E))
print("FAILURES:", len(F))
