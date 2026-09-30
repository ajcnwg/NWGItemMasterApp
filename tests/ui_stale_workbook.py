"""A workbook downloaded, then the app changes, then the old download is uploaded."""
from testbase import BASE, IMPORT_SNAP, KEEP
import io
from uiharness import *
import openpyxl
import streamlit as st_mod
from itemmaster import dept_mapping as dm, old_workbook_import as owi
from itemmaster.db import get_engine
from sqlalchemy import text
E = get_engine(); F = []
BLANK = BASE


def check(ok, msg):
    print(("  PASS " if ok else "  FAIL ") + msg)
    if not ok: F.append(msg)


def rows(ws):
    hdr = [c.value for c in ws[2]]
    return hdr, [(i, dict(zip(hdr, [c.value for c in r]))) for i, r in enumerate(ws.iter_rows(min_row=3), start=3)]


def setc(ws, hdr, i, col, v):
    ws.cell(row=i, column=hdr.index(col) + 1).value = v


def state(cid):
    with E.connect() as c:
        return tuple(c.execute(text("SELECT decision_state, decided_department FROM dbo.dept_mapping_combos WHERE combo_id=:c"), {"c": cid}).one())


key = lambda r: tuple(owi._s(r.get(c)).upper() for c in ("Source", "Old Department", "Category", "Subcategory"))
try:
    dm.restore_snapshot(E, BLANK, "AJ")
    src = dict(E.connect().execute(text("SELECT source_key, source_label FROM dbo.sources")).all())
    gk = lambda r: (str(src.get(r["source_key"]) or r["source_key"]).upper(), *(owi._s(r[c]).upper() for c in ("raw_department", "raw_category", "raw_subcategory")))
    q = dm.get_review_queue(E, "review")
    q = q[q["suggested_department"].notna() & q["n_upcs_total"].between(2, 30)].sort_values("combo_id").to_dict("records")
    A, B, E_ = q[0], q[1], q[2]            # A: Jason stages it; B: Jason breaks it out; E_: untouched
    dec = dm.get_decided_combos(E)
    dec = dec[dec["status"].str.startswith("Whole") & dec["n_upcs_total"].between(2, 30)].sort_values("combo_id").to_dict("records")
    D, Fg = dec[0], dec[1]                 # D: someone pushes a decision on it; Fg: untouched, sent to Broken Out
    bo = dm.get_broken_out_combos(E).sort_values("combo_id")
    C = int(bo.iloc[0]["combo_id"])
    items = dm.get_pending_upc_overrides(E, C)
    u1, u2 = items["upc"].iloc[0], items["upc"].iloc[1]   # u1: Jason stages it; u2: untouched
    print("  groups: A", A["combo_id"], "B", B["combo_id"], "E", E_["combo_id"], "D", D["combo_id"], "F", Fg["combo_id"], "| BO", C, u1, u2)

    print("\n== 1. AJ downloads")
    data = owi.export_workbook(E)

    print("\n== 2. Meanwhile, in the app")
    dm.upsert_combo_suggestion(E, int(A["combo_id"]), "review", "DELI", A["source_key"], "x", int(A["n_upcs_total"]), "Jason")
    dm.break_out_combo(E, int(B["combo_id"]), "Jason")
    dm.approve_combo(E, int(D["combo_id"]), "PRODUCE", "Kristi", pushed_by="AJ")
    dm.stage_broken_out_decisions(E, {u1: {"department": "FROZEN", "combo_id": C, "label": "x", "description": "", "source_key": ""}}, "Jason")
    print("   Jason staged A → DELI, broke out B; Kristi's D → PRODUCE was pushed; Jason staged item", u1, "→ FROZEN")

    print("\n== 3. AJ edits his old download")
    wb = openpyxl.load_workbook(io.BytesIO(data))
    for sheet in ("Department Mapping Crosswalk", "Department Mapping Unmatched"):
        ws = wb[sheet]; hdr, rs = rows(ws)
        for i, r in rs:
            for g, d in ((A, "BAKERY"), (B, "BAKERY"), (E_, "BAKERY")):
                if key(r) == gk(g):
                    setc(ws, hdr, i, "Action", "Approve"); setc(ws, hdr, i, "Manual Override Department", d)
    ws = wb["Department Mapping Decided"]; hdr, rs = rows(ws)
    for i, r in rs:
        if key(r) == gk(D):
            setc(ws, hdr, i, "Manual Override Department", "MEAT")
        if key(r) == gk(Fg):
            setc(ws, hdr, i, "Action", "Send to Broken Out")
    ws = wb["Department UPC Overrides"]; hdr, rs = rows(ws)
    for i, r in rs:
        if owi.clean_upc(r["UPC"]) in (u1, u2):
            setc(ws, hdr, i, "Manual Override Department", "BAKERY")
    buf = io.BytesIO(); wb.save(buf); edited = buf.getvalue()

    print("\n== 4. AJ uploads it")
    real = st_mod.file_uploader
    st_mod.file_uploader = lambda label, *a, key=None, **k: (type("U", (io.BytesIO,), {"name": "old download.xlsx"})(edited)
                                                             if key and key.startswith("owi_file_") else real(label, *a, key=key, **k))
    st_mod.cache_data.clear()
    a = session("aj"); run(a, "AJ load"); goto(a, "Department Review", "Settings")
    m = {x.label: x.value for x in a.metric}
    warn = [x.value for x in a.warning if "skipped" in x.value]
    print("   report:", m)
    print("   ", warn[:1])
    p = a.session_state["_owi"]["plan"]
    conf = {int(r["combo_id"]) for df in (p["groups"], p["moves"]) for r in df.to_dict("records") if str(r["Result"]).startswith(owi.CONFLICT)}
    conf_items = set(p["items"].loc[p["items"]["Result"].str.startswith(owi.CONFLICT), "UPC"])
    check(conf == {int(D["combo_id"])} and conf_items == {u1}, f"D and item u1 are skipped as changed-since-download ({conf}, {conf_items})")
    check(bool(warn) and "2 change(s)" in warn[0], "a clear warning says so, and why")
    res = dict(zip(p["groups"]["combo_id"].astype(int), p["groups"]["Result"]))
    print("   A:", res.get(int(A["combo_id"]))); print("   B:", res.get(int(B["combo_id"])))
    check(str(res.get(int(A["combo_id"]))).startswith("Skipped — Jason already has DELI staged"), "A skipped: Jason already has DELI staged")
    check(str(res.get(int(B["combo_id"]))).startswith("Skipped — it's in Broken Out in the app now"), "B skipped: it's in Broken Out now (not silent)")
    print("   e.g.", next(r for r in p["groups"]["Result"] if str(r).startswith(owi.CONFLICT)))
    a.button(key="owi_apply").click(); run(a, "Stage")
    st_mod.file_uploader = real
    pc = dm.get_pending_changes(E)
    check(pc.get(int(A["combo_id"]), {}).get("department") == "DELI" and pc[int(A["combo_id"])]["staged_by"] == "Jason", "Jason's A → DELI is untouched")
    check(state(int(B["combo_id"]))[0] == "broken_out", "B stays broken out (Jason's move)")
    check(state(int(D["combo_id"])) == ("decided", "PRODUCE") and int(D["combo_id"]) not in pc, "Kristi's pushed D → PRODUCE is untouched")
    pu = dm.get_pending_upc_changes(E)
    check(pu.get(u1, {}).get("department") == "FROZEN" and pu[u1]["staged_by"] == "Jason", "Jason's item u1 → FROZEN is untouched")
    check(pc.get(int(E_["combo_id"]), {}).get("department") == "BAKERY", "the untouched group E was staged")
    check(pu.get(u2, {}).get("department") == "BAKERY", "the untouched item u2 was staged")
    check(state(int(Fg["combo_id"]))[0] == "broken_out", "the untouched group F was sent to Broken Out")

    print("\n== 5. Top-bar Undo / Redo on what the upload did")
    a = session("aj"); run(a, "AJ load"); goto(a, "Department Review", "Pending Changes")
    a.button(key="topbar_undo").click(); run(a, "Undo")
    next(b for b in a.button if (b.label or "").startswith("Confirm")).click(); run(a, "confirm")
    after_undo = dm.get_pending_upc_changes(E).get(u2)
    check(after_undo is None, "Undo took back the upload's last step (item u2)")
    a = session("aj"); run(a, "AJ load"); goto(a, "Department Review", "Pending Changes")
    a.button(key="topbar_redo").click(); run(a, "Redo")
    b = [x for x in a.button if (x.label or "").startswith("Confirm")]
    if b:
        b[0].click(); run(a, "confirm")
    check(dm.get_pending_upc_changes(E).get(u2, {}).get("department") == "BAKERY", "Redo put it back")
    check(dm.get_pending_upc_changes(E).get(u1, {}).get("staged_by") == "Jason", "…and Jason's u1 is still his")
    check(len(a.exception) == 0 and not any("Something went wrong" in x.value for x in a.markdown), "no errors")
finally:
    print("\n  back to the blank baseline…"); dm.restore_snapshot(E, BLANK, "AJ")
    for s_ in dm.list_snapshots(E).to_dict("records"):
        if s_["snapshot_id"] not in KEEP:
            dm.delete_snapshot(E, s_["snapshot_id"])
    with E.begin() as c:
        c.execute(text("DELETE FROM dbo.user_workspace")); c.execute(text("DELETE FROM dbo.change_discard_notices"))
    print("  vs #%d:" % BLANK, dm.compare_snapshot_to_live(E, BLANK), dm.old_workbook_import_summary(E))
print("FAILURES:", len(F))
