import os
"""The whole workbook process on the blank baseline: download, edit, upload —
then download again, clear some of your own staged rows (and someone else's),
upload, and check only yours are taken back."""
from testbase import BASE, IMPORT_SNAP, KEEP
import io
from uiharness import *
import openpyxl
import streamlit as st_mod
from itemmaster import dept_mapping as dm, old_workbook_import as owi
from itemmaster.db import get_engine
from sqlalchemy import text
E = get_engine(); F = []
BLANK = BASE
SP = os.path.dirname(os.path.abspath(__file__))
OLD = __import__('testbase').V2


def check(ok, msg):
    print(("  PASS " if ok else "  FAIL ") + msg)
    if not ok: F.append(msg)


def counts():
    s = dm.old_workbook_import_summary(E)
    with E.connect() as c:
        allp = c.execute(text("SELECT (SELECT COUNT(*) FROM dbo.dept_mapping_pending_changes), (SELECT COUNT(*) FROM dbo.dept_mapping_pending_upc_changes), "
                              "(SELECT COUNT(*) FROM dbo.item_master_pending_changes)")).one()
    return tuple(allp), s


def upload(who, data, name, apply=True):
    real = st_mod.file_uploader
    st_mod.file_uploader = lambda label, *a, key=None, **k: (type("U", (io.BytesIO,), {"name": name})(data)
                                                             if key and key.startswith("owi_file_") else real(label, *a, key=key, **k))
    try:
        st_mod.cache_data.clear()
        at = session(who); run(at, f"{who} load"); goto(at, "Department Review", "Settings")
        m = {x.label: x.value for x in at.metric}
        print("   report:", m)
        if apply:
            at.button(key="owi_apply").click(); run(at, "Stage")
            print("   ", [x.value[:260] for x in at.success if "Imported" in x.value][:1])
        return at, m
    finally:
        st_mod.file_uploader = real


def rows(ws):
    hdr = [c.value for c in ws[2]]
    return hdr, [(i, dict(zip(hdr, [c.value for c in r]))) for i, r in enumerate(ws.iter_rows(min_row=3), start=3)]


def setc(ws, hdr, i, col, v):
    ws.cell(row=i, column=hdr.index(col) + 1).value = v  # (cell(..., value=None) would leave it unchanged)


key = lambda r: tuple(owi._s(r.get(c)).upper() for c in ("Source", "Old Department", "Category", "Subcategory"))
try:
    dm.restore_snapshot(E, BLANK, "AJ")
    check(counts() == ((0, 0, 0), {"groups": 0, "items": 0, "overrides": 0, "moves": 0}), f"blank baseline: nothing staged {counts()}")

    print("\n== 1. Download the app as a workbook")
    a = session("aj"); run(a, "admin load"); goto(a, "Department Review", "Settings")
    check(any(d.proto.label.strip() == "Download the app as a workbook" for d in a.get("download_button")), "Download button in Settings")
    data = owi.export_workbook(E)  # what that button builds
    open(SP + r"\1 - downloaded from the app (blank baseline).xlsx", "wb").write(data)
    wb = openpyxl.load_workbook(io.BytesIO(data))

    print("\n== 2. Make the changes in Excel (the old workbook's hand-made decisions)")
    ex = owi.extract(owi.read_workbook(OLD))
    want_groups = {g["key"]: g["department"] for g in ex["groups"]}
    edits = {"approve": 0, "decided": 0, "moves": 0, "items": 0}
    for sheet in ("Department Mapping Crosswalk", "Department Mapping Unmatched"):
        ws = wb[sheet]; hdr, rs = rows(ws)
        for i, r in rs:
            if key(r) in want_groups:
                setc(ws, hdr, i, "Action", "Approve"); setc(ws, hdr, i, "Manual Override Department", want_groups[key(r)])
                edits["approve"] += 1
    ws = wb["Department Mapping Decided"]; hdr, rs = rows(ws)
    moves = {m["key"]: m for m in ex["moves"]}
    for i, r in rs:
        if key(r) in want_groups:
            setc(ws, hdr, i, "Manual Override Department", want_groups[key(r)]); edits["decided"] += 1
        elif key(r) in moves and moves[key(r)]["move"] == "to_broken_out":
            setc(ws, hdr, i, "Action", "Send to Broken Out"); edits["moves"] += 1
    ws = wb["Department Mapping Broken Out"]; hdr, rs = rows(ws)
    supplies_key = next(k for k in want_groups if "SUPPLIES NOT FOR RESALE" in k)
    for i, r in rs:
        if key(r) in moves and moves[key(r)]["move"] == "to_review":
            setc(ws, hdr, i, "Action", "Return to Crosswalk/Unmatched"); edits["moves"] += 1
    item_want = {it["upc"]: it["department"] for it in ex["items"]}
    ws = wb["Department UPC Overrides"]; hdr, rs = rows(ws)
    for i, r in rs:
        u = owi.clean_upc(r["UPC"])
        d = item_want.get(u) or (want_groups[supplies_key] if key(r) == supplies_key else None)
        if d:
            setc(ws, hdr, i, "Manual Override Department", d); edits["items"] += 1
    print("   edits made:", edits)
    buf = io.BytesIO(); wb.save(buf); edited = buf.getvalue()
    open(SP + r"\2 - edited in Excel.xlsx", "wb").write(edited)

    print("\n== 3. Upload it in Settings")
    a, m = upload("aj", edited, "2 - edited in Excel.xlsx")
    n, s = counts()
    print("   staged:", n, s)
    check(m.get("To take back") == "0" and m.get("Skipped") == "0", "report: nothing to take back or skip")
    check(s == {"groups": 22, "items": 175, "overrides": 0, "moves": 3},
          f"staged exactly the edits: 22 group decisions, 175 Broken Out item decisions (STORE SUPPLIES 116 + SS COOKIES 59), 3 moves, 0 UPC overrides {s}")
    res = [x.value for x in a.success if "Imported" in x.value]
    check(bool(res), "done message with a full report download")

    print("\n== 4. Download again, clear some of my own staged rows (and one of Jason's), upload")
    j = session("jason"); run(j, "Jason signs in")
    q = dm.get_review_queue(E, "review")
    busy = set(dm.get_pending_changes(E)) | {m_["combo_id"] for m_ in dm.get_recent_moves(E)}
    jr = q[q["suggested_department"].notna() & ~q["combo_id"].isin(busy) & q["n_upcs_total"].between(2, 30)].iloc[0]
    j.session_state["dept_shared_filter"] = {"search": max([jr.raw_category or "", jr.raw_subcategory or ""], key=len), "sort_label": None, "sort_desc": True, "page_size": 25}
    goto(j, "Department Review", "Crosswalk")
    j.selectbox(key=f"dept_choice_review_{jr.combo_id}").select(jr.suggested_department); run(j, "pick")
    j.button(key=f"approve_review_{jr.combo_id}").click(); run(j, "Jason approves a group")
    check(int(jr.combo_id) in dm.get_pending_changes(E), "Jason has his own staged approval")
    data2 = owi.export_workbook(E)
    wb2 = openpyxl.load_workbook(io.BytesIO(data2))
    ws = wb2["Department Mapping Crosswalk"]; hdr, rs = rows(ws)
    cleared_groups, jason_row = [], None
    for i, r in rs:
        if r["Action"] == "Approve":
            if key(r)[3] == (jr.raw_subcategory or "").upper() and key(r)[2] == (jr.raw_category or "").upper():
                setc(ws, hdr, i, "Action", "Not Yet Reviewed"); setc(ws, hdr, i, "Manual Override Department", None); jason_row = i
            elif len(cleared_groups) < 2 and "BULK" in key(r):
                setc(ws, hdr, i, "Action", "Not Yet Reviewed"); setc(ws, hdr, i, "Manual Override Department", None)
                cleared_groups.append(key(r))
    ws = wb2["Department UPC Overrides"]; hdr, rs = rows(ws)
    cleared_items = []
    for i, r in rs:
        if r["Decided Via"] == "Staged (in Pending Changes)" and "COOKIES" in (r["Subcategory"] or "") and len(cleared_items) < 5:
            setc(ws, hdr, i, "Manual Override Department", None); cleared_items.append(owi.clean_upc(r["UPC"]))
    print("   cleared: 2 of my Crosswalk approvals", [k[3] for k in cleared_groups], "+ 5 of my SS COOKIES items + Jason's approval (row", jason_row, ")")
    buf = io.BytesIO(); wb2.save(buf); edited2 = buf.getvalue()
    open(SP + r"\3 - downloaded again, some staged rows cleared.xlsx", "wb").write(edited2)
    before, _ = counts()
    a, m = upload("aj", edited2, "3 - cleared.xlsx")
    check(m.get("To take back") == "7" and m.get("To stage") == "0", f"report: 7 to take back (2 groups + 5 items), nothing new to stage ({m})")
    after, s = counts()
    check(after == (before[0] - 2, before[1] - 5, before[2]), f"taken back: 2 group decisions and 5 item decisions {before} -> {after}")
    check(int(jr.combo_id) in dm.get_pending_changes(E), "Jason's staged approval was NOT taken back (only your own changes)")
    tb = a.session_state["_owi_done"]["plan"]["take_back"]
    check(any("staged by Jason" in r for r in tb["Result"]), "…and the report says so (kept — staged by Jason)")

    print("\n== 5. Top-bar Undo brings a take-back back")
    a = session("aj"); run(a, "AJ load"); goto(a, "Department Review", "Pending Changes")
    a.button(key="topbar_undo").click(); run(a, "top-bar Undo")
    next(b for b in a.button if (b.label or "").startswith("Confirm")).click(); run(a, "confirm")
    n2, _ = counts()
    check(n2[1] == after[1] + 5 or n2[0] == after[0] + 1, f"the last take-back is staged again {after} -> {n2}")
    check(len(a.exception) == 0 and not any("Something went wrong" in x.value for x in a.markdown), "no errors")
finally:
    print("\n  back to the blank baseline…"); dm.restore_snapshot(E, BLANK, "AJ")
    with E.begin() as c:
        c.execute(text("DELETE FROM dbo.user_workspace")); c.execute(text("DELETE FROM dbo.change_discard_notices"))
    for s_ in dm.list_snapshots(E).to_dict("records"):
        if s_["snapshot_id"] not in KEEP:
            dm.delete_snapshot(E, s_["snapshot_id"])
    print("  vs #%d:" % BLANK, dm.compare_snapshot_to_live(E, BLANK), counts(), "snapshots:", dm.list_snapshots(E)["snapshot_id"].tolist())
print("FAILURES:", len(F))
