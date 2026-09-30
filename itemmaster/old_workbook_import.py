"""
Admin-only: bring the decisions PEOPLE made in an old department workbook
(the Excel process from before this app — "..._department_workbook.xlsx")
into the app. Nothing auto-applied is taken, only what someone decided:

  * Crosswalk / Unmatched groups someone approved (their Manual Override
    Department if they typed one, otherwise New Department);
  * Decided groups someone manually overrode;
  * items someone manually reviewed or approved (Department UPC Overrides,
    Decided UPC Overrides) and anything on Final UPC Overrides;
  * the moves someone asked for (Break Out, Return to Crosswalk / Unmatched,
    Send to Broken Out, Move to Undecided).

Each one lands where it belongs in the app: an app group whose items ALL
get one Department gets a group decision, an item in a Broken Out group
gets that group's item decision, anything else gets a UPC override of just
its Department. Anything the app ALREADY does — the group already decided
that way, the item already has that Department, the move already made, or
the same thing already staged — is skipped and shown as "Already done".
Everything else is staged to Pending Changes (moves happen straight away,
as they do in the app), and every line is in the report.
"""

import io
import json
from datetime import datetime, timezone
from collections import defaultdict

import openpyxl
import pandas as pd
import re

from sqlalchemy import bindparam, text

from itemmaster import dept_mapping as dm
from itemmaster.ingest import clean_upc

SOURCE_KEYS = {"KEHE": "kehe", "SPINS": "spins", "C&S PNW": "cs_pnw", "C&S CA": "cs_ca", "URM": "urm",
               "SCAN ADVANTAGE": "nwg", "NWG": "nwg"}
SHEETS = ["Department Mapping Crosswalk", "Department Mapping Unmatched", "Department Mapping Broken Out",
          "Department Mapping Decided", "Decided Broken Out Combos", "Department UPC Overrides",
          "Decided UPC Overrides", "Final UPC Overrides", "Department Mapping Detail", "App staged (do not edit)",
          "App state (do not edit)"]
STAGED_SHEET = "App staged (do not edit)"  # written by export_workbook: what was already staged, and by whom
STATE_SHEET = "App state (do not edit)"    # written by export_workbook: every group's / item's state at download time
CONFLICT = "Skipped — changed in the app since this workbook was downloaded"
KEY = ["Source", "Old Department", "Category", "Subcategory"]
DONE = "Already done"
NOTE_PREFIX = "📥 One-time old-workbook import"
STAGE = "Will be staged"
MOVE = "Will be moved now"
TAKE = "Will be taken back"
SKIP = "Skipped"
WAIT = "Waiting on your choice"  # in Needs your choice: nothing staged until someone picks


def _s(v) -> str:
    return "" if v is None or (isinstance(v, float) and pd.isna(v)) else str(v).strip()


def read_workbook(file) -> dict:
    """The sheets this needs, as text DataFrames (headers are on row 2 of
    every sheet). Missing sheets come back empty."""
    wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
    out = {}
    for name in SHEETS:
        if name not in wb.sheetnames:
            out[name] = pd.DataFrame()
            continue
        rows = wb[name].iter_rows(min_row=2, values_only=True)
        header = next(rows, None) or ()
        width = max((i + 1 for i, h in enumerate(header) if h is not None), default=0)
        cols = [_s(h) for h in header[:width]]
        data = [[_s(v) for v in r[:width]] for r in rows if any(v not in (None, "") for v in r[:width])]
        out[name] = pd.DataFrame(data, columns=cols)
    wb.close()
    if all(out[n].empty for n in SHEETS[:8]):
        raise ValueError("This doesn't look like an old department workbook — none of its Department Mapping sheets were found.")
    return out


def _key(r) -> tuple:
    return tuple(_s(r.get(k)).upper() for k in KEY)


def _label(key) -> str:
    src, *bits = key
    bits = [b for b in bits if b]
    return f"{src} — {' / '.join(bits) if bits else '(blank)'}"


def extract(sheets: dict) -> dict:
    """Only what a person decided. Returns {"groups": [...], "items": [...],
    "moves": [...]} in the old workbook's own terms."""
    groups, items, moves = [], [], []
    for sheet, area in (("Department Mapping Crosswalk", "Crosswalk"), ("Department Mapping Unmatched", "Unmatched")):
        for r in sheets[sheet].to_dict("records"):
            action = _s(r.get("Action"))
            if action == "Approve" or _s(r.get("Approved?")) == "Yes":
                dept = _s(r.get("Manual Override Department")) or _s(r.get("New Department"))
                if dept:
                    groups.append({"key": _key(r), "department": dept, "from": f"Old {area}: approved group"})
            elif action.startswith("Break Out") or _s(r.get("Break Out to UPC-Level?")) == "Yes":
                moves.append({"key": _key(r), "move": "to_broken_out", "from": f"Old {area}: Break Out to UPC-Level"})
    for r in sheets["Department Mapping Decided"].to_dict("records"):
        action, manual = _s(r.get("Action")), _s(r.get("Manual Override Department"))
        if manual and action in ("", "Keep (No Change)", "Manually Overridden"):
            groups.append({"key": _key(r), "department": manual, "from": "Old Decided: manually overridden group"})
        elif action == "Send to Broken Out":
            moves.append({"key": _key(r), "move": "to_broken_out", "from": "Old Decided: Send to Broken Out"})
        elif action.startswith("Return to"):
            moves.append({"key": _key(r), "move": "to_review", "from": f"Old Decided: {action}"})
    for r in sheets["Department Mapping Broken Out"].to_dict("records"):
        if _s(r.get("Action")).startswith("Return to"):
            moves.append({"key": _key(r), "move": "to_review", "from": "Old Broken Out: Return to Crosswalk/Unmatched"})
    for r in sheets["Decided Broken Out Combos"].to_dict("records"):
        action, manual = _s(r.get("Action")), _s(r.get("Manual Override Department"))
        if manual and action in ("", "Keep (No Change)"):
            # decide the whole group (replaces its item decisions when pushed)
            groups.append({"key": _key(r), "department": manual, "from": "Old Decided Broken Out: manual override"})
        elif action.startswith("Return to"):
            moves.append({"key": _key(r), "move": "to_review", "from": f"Old Decided Broken Out: {action}"})
        elif action.startswith("Move to Undecided"):
            moves.append({"key": _key(r), "move": "reopen", "from": "Old Decided Broken Out: Move to Undecided"})
        elif action.startswith("Move to Decided Combos") and (manual or _s(r.get("New Department"))):
            groups.append({"key": _key(r), "department": manual or _s(r.get("New Department")),
                           "from": "Old Decided Broken Out: Move to Decided Combos"})
    for sheet in ("Department UPC Overrides", "Decided UPC Overrides"):
        for r in sheets[sheet].to_dict("records"):
            via, manual = _s(r.get("Decided Via")), _s(r.get("Manual Override Department"))
            if via.startswith("Manually") or manual:
                dept = manual or _s(r.get("Final Department")) or _s(r.get("New Department"))
                if dept:
                    items.append({"upc": clean_upc(r.get("UPC")), "department": dept,
                                  "from": f"Old {sheet}: {via or 'manual override'}", "rank": 2})
    for r in sheets["Final UPC Overrides"].to_dict("records"):
        if _s(r.get("UPC")) and _s(r.get("New Department")):
            items.append({"upc": clean_upc(r.get("UPC")), "department": _s(r.get("New Department")),
                          "from": "Old Final UPC Overrides", "rank": 3})
    # A group's items, from the old workbook's own Detail sheet.
    detail = sheets["Department Mapping Detail"]
    members = defaultdict(list)
    if not detail.empty:
        for r in detail.to_dict("records"):
            members[_key(r)].append(clean_upc(r.get("UPC")))
    for g in groups:
        g["upcs"] = members.get(g["key"], [])
    staged = []
    for r in sheets.get(STAGED_SHEET, pd.DataFrame()).to_dict("records"):
        staged.append({"kind": _s(r.get("Kind")), "key": _key(r), "upc": clean_upc(r.get("UPC")) if _s(r.get("UPC")) else "",
                       "department": _s(r.get("Department")), "staged_by": _s(r.get("Staged by"))})
    state = {"downloaded_at": None, "groups": {}, "items": {}}
    for r in sheets.get(STATE_SHEET, pd.DataFrame()).to_dict("records"):
        if _s(r.get("Kind")) == "downloaded":
            state["downloaded_at"] = _s(r.get("State"))
        elif _s(r.get("Kind")) == "group":
            state["groups"][_key(r)] = _s(r.get("State"))
        elif _s(r.get("Kind")) == "item":
            state["items"][clean_upc(r.get("UPC"))] = _s(r.get("State"))
    placements = defaultdict(list)  # key -> [(sheet place, instruction)] for every row of the placement sheets
    for sheet, place in PLACEMENT_SHEETS.items():
        for r in sheets[sheet].to_dict("records"):
            action, manual = _s(r.get("Action")), _s(r.get("Manual Override Department"))
            if _s(r.get("Approved?")) == "Yes" and not action:
                action = "Approve"
            what = action if action and action not in KEEP_ACTIONS else ""
            if manual and (not what or what in ("Approve", "Manually Overridden")):
                what = f"{'Approve as' if what == 'Approve' else 'Change to'} {manual}"
            elif what == "Approve":
                what = f"Approve as {_s(r.get('New Department'))}"
            placements[_key(r)].append({"place": place, "what": what or "Keep (no change)",
                                        "items": int(float(_s(r.get("Total UPCs")) or 0)),
                                        "spelled": _s(r.get("Old Department")), "now": _s(r.get("New Department"))})
    return {"groups": groups, "items": items, "moves": moves, "staged": staged, "state": state,
            "placements": dict(placements)}


PLACEMENT_SHEETS = {"Department Mapping Crosswalk": "Crosswalk", "Department Mapping Unmatched": "Unmatched",
                    "Department Mapping Broken Out": "Broken Out", "Department Mapping Decided": "Decided",
                    "Decided Broken Out Combos": "Decided item by item"}
KEEP_ACTIONS = {"", "Keep (No Change)", "Keep (Broken Out)", "Not Yet Reviewed", "Keep"}


def _sheet_of(from_: str) -> str:
    """ "Old Decided Broken Out: manual override" -> "Decided item by item"."""
    for prefix, place in (("Old Decided Broken Out", "Decided item by item"), ("Old Crosswalk", "Crosswalk"),
                          ("Old Unmatched", "Unmatched"), ("Old Broken Out", "Broken Out"), ("Old Decided", "Decided")):
        if from_.startswith(prefix):
            return place
    return ""


def _place(state: str, tier, n_evidence, decided) -> str:
    """Where a group sits in the app, in the workbook's terms."""
    if state == "broken_out":
        return "Broken Out"
    if state == "decided_broken_out":
        return "Decided item by item"
    if state == "decided" or decided:
        return "Decided"
    return "Unmatched" if dm.origin_tier(tier, n_evidence) == "unmatched" else "Crosswalk"


def _same_place(a: str, b: str) -> bool:
    queue = {"Crosswalk", "Unmatched"}
    return a == b or (a in queue and b in queue)


def _group_state(state, decided, pend) -> str:
    return "|".join(_s(v) for v in (state, decided, *(pend or ("", ""))))


def _item_state(dept, via, staged) -> str:
    return "|".join(_s(v) for v in (dept, via, *(staged or ("", ""))))


def _describe_state(sig: str) -> str:
    """"decided|BAKERY|GROCERY|Jason" -> plain words for the report."""
    bits = (sig.split("|") + ["", "", "", ""])[:4]
    state, decided, pend, who = bits
    where = {"broken_out": "in Broken Out", "decided_broken_out": "finished item by item", "decided": f"decided as {decided}"}.get(
        state, (f"auto-decided as {decided}" if decided else "undecided") if state in ("not_reviewed", "") else state)
    return where + (f", {who} has {pend} staged" if pend else "")


def _app_groups(engine) -> pd.DataFrame:
    return pd.read_sql(text(
        "SELECT combo_id, source_key, raw_department, raw_category, raw_subcategory, tier, n_evidence, "
        "decision_state, decided_department, manual_department, n_upcs_total FROM dbo.dept_mapping_combos"), engine)


def _where(state: str, tier, n_evidence) -> str:
    if state in ("broken_out",):
        return "Broken Out"
    if state in ("decided", "decided_broken_out"):
        return "Decided"
    return "Unmatched" if dm.origin_tier(tier, n_evidence) == "unmatched" else "Crosswalk"


def plan(engine, extracted: dict, departments: list, actor: str) -> dict:
    """Works out, line by line, what importing would do — nothing is written."""
    canon = {d.upper(): d for d in departments}
    combos = _app_groups(engine)
    combos["label"] = combos.apply(lambda c: f"{(c.source_key or '').upper()} — " + (
        " / ".join(b for b in (c.raw_department, c.raw_category, c.raw_subcategory) if b) or "(blank)"), axis=1)
    by_key = {}
    records = [{**c, "decided_department": _s(c["decided_department"]) or None,  # database NULLs, not NaN
                "manual_department": _s(c["manual_department"]) or None} for c in combos.to_dict("records")]
    for c in records:
        by_key[(c["source_key"], _s(c["raw_department"]).upper(), _s(c["raw_category"]).upper(),
                _s(c["raw_subcategory"]).upper())] = c
    state = {c: s for c, s in zip(combos.combo_id, combos.decision_state)}
    info = {c["combo_id"]: c for c in records}
    app_place = {c["combo_id"]: _place(c["decision_state"], c["tier"], c["n_evidence"], c["decided_department"]) for c in records}

    # the app's own source names first (so a source added later works too), then the old file's
    labels = {str(lb).strip().upper(): k for k, lb in pd.read_sql(text("SELECT source_key, source_label FROM dbo.sources"),
                                                                  engine).itertuples(index=False) if lb}

    def app_group(old_key):
        src, *rest = old_key
        return by_key.get((labels.get(src) or SOURCE_KEYS.get(src, src.lower()), *rest))

    with engine.connect() as conn:
        mem = pd.read_sql(text("SELECT upc, combo_id FROM dbo.dept_mapping_combo_upcs"), conn)
        pend_groups = {r[0]: (r[1], r[2]) for r in conn.execute(text(
            "SELECT combo_id, department, staged_by FROM dbo.dept_mapping_pending_changes")).all()}
    combo_of = dict(zip(mem.upc, mem.combo_id))

    # 1. Moves. A Broken Out group whose items a person decided in the workbook
    #    stays in Broken Out even if the workbook also says to return it — its
    #    item decisions are staged there instead of being lost.
    decided_item_groups = {combo_of.get(it["upc"]) for it in extracted["items"] if it["rank"] == 2}
    move_rows = []
    for m in extracted["moves"]:
        c = app_group(m["key"])
        row = {"Group (old workbook)": _label(m["key"]), "Asked for": m["from"], "App group": "", "Now": "",
               "Result": "", "combo_id": None, "move": m["move"]}
        if c is None:
            row.update(Result=f"{SKIP} — no matching group in the app")
        else:
            s = state[c["combo_id"]]
            row.update({"App group": c["label"], "Now": _where(s, c["tier"], c["n_evidence"]), "combo_id": c["combo_id"],
                        "n_upcs_total": int(c["n_upcs_total"] or 0)})
            done = ((m["move"] == "to_broken_out" and s in ("broken_out", "decided_broken_out"))
                    or (m["move"] == "to_review" and s == "not_reviewed" and not c["decided_department"])
                    or (m["move"] == "reopen" and s == "broken_out"))
            if m["move"] == "reopen" and s != "decided_broken_out" and not done:
                row["Result"] = f"{SKIP} — it isn't a finished Broken Out group in the app"
            elif m["move"] == "to_review" and s in ("broken_out", "decided_broken_out") and c["combo_id"] in decided_item_groups:
                row["Result"] = (f"{SKIP} — kept in Broken Out: a person decided some of its items in the workbook, "
                                 "and those decisions are staged there")
            elif done:
                row["Result"] = f"{DONE} — already {row['Now']}"
            else:
                row["Result"] = MOVE
                state[c["combo_id"]] = {"to_broken_out": "broken_out", "to_review": "not_reviewed",
                                        "reopen": "broken_out"}[m["move"]]
        move_rows.append(row)

    # 2. Every item's wanted Department (item-level beats group-level, the last sheet wins ties)
    want, why = {}, {}
    for g in extracted["groups"]:
        for u in g["upcs"]:
            if u not in want or why[u][1] < 1:
                want[u], why[u] = g["department"], (g["from"], 1)
    for it in sorted(extracted["items"], key=lambda x: x["rank"]):
        want[it["upc"]], why[it["upc"]] = it["department"], (it["from"], it["rank"])
    bad_dept = {u for u, d in want.items() if d.upper() not in canon}
    want = {u: canon.get(d.upper(), d) for u, d in want.items()}

    # 3. Group decisions. (a) A group someone approved in the workbook becomes
    #    a decision on the app's group of the same name — even when the app's
    #    group holds items the workbook's didn't (those are counted as
    #    "Also covers"). (b) Any other app group whose items ALL got the same
    #    Department from the workbook gets a group decision too.
    group_rows, whole = [], set()
    sizes = mem.groupby("combo_id").size()
    members_of = mem.groupby("combo_id")["upc"].apply(set).to_dict()
    by_cid = {}
    for g in extracted["groups"]:
        c = app_group(g["key"])
        if c is not None:
            by_cid.setdefault(c["combo_id"], []).append(g)

    def group_row(cid, d, from_, wb_items, also):
        s = state[cid]
        c = info[cid]
        now = c["decided_department"] if s == "decided" or c["decided_department"] else None
        if s == "decided_broken_out":
            now = "decided item by item"
        pend = pend_groups.get(cid)
        row = {"Area": _where(s, c["tier"], c["n_evidence"]), "Group": c["label"], "Items": int(sizes.get(cid, 0)),
               "Workbook items": wb_items, "Also covers": also,
               "Now": now or "(undecided)", "Will be": d, "From (old workbook)": from_,
               "combo_id": int(cid), "tier": c["tier"], "n_evidence": c["n_evidence"],
               "source_key": c["source_key"], "n_upcs_total": int(c["n_upcs_total"])}
        if now == d:
            how = "by a person" if (s == "decided" or c["manual_department"]) else "automatically"
            row["Result"] = f"{DONE} — already decided as {d} ({how})"
        elif pend and pend[0] == d:
            row["Result"] = f"{DONE} — already staged ({pend[1]})"
        elif pend and pend[1] != actor:
            row["Result"] = f"{SKIP} — {pend[1]} already has {pend[0]} staged for this group"
        else:
            row["Result"] = STAGE
        return row

    for cid, gs in by_cid.items():
        # A group someone approved as a whole gets a whole-group decision — even
        # one the app has since finished item by item (Decided): the decision
        # replaces its item decisions when pushed. Only a group still being
        # worked on in Broken Out keeps its items decided one by one (step 4).
        if state[cid] not in ("not_reviewed", "decided", "decided_broken_out"):
            # Being decided item by item in the app: a whole-group decision
            # can't go on it — say so (its items, if the workbook lists them,
            # are decided one by one in step 4).
            members = members_of.get(cid, set())
            asked = {u for g in gs for u in g["upcs"]} & members
            row = group_row(cid, gs[0]["department"], gs[0]["from"], len(asked), 0)
            row["Result"] = (f"{SKIP} — it's in Broken Out in the app now, being decided item by item"
                             + (f"; its {len(asked):,} item(s) are staged one by one instead" if asked else
                                "; set its items on the Department UPC Overrides sheet instead"))
            group_rows.append(row)
            continue
        votes = {}
        for g in gs:
            dd = canon.get(g["department"].upper())
            if dd:
                votes[dd] = votes.get(dd, 0) + len(g["upcs"]) + 1
        if not votes:
            continue
        d = max(votes, key=votes.get)
        members = members_of.get(cid, set())
        wb_upcs = {u for g in gs for u in g["upcs"]}
        # an item with its own, different, person decision keeps it (step 4)
        whole |= {u for u in members if want.get(u, d) == d}
        from_ = next(g["from"] for g in gs if canon.get(g["department"].upper()) == d)
        group_rows.append(group_row(cid, d, from_, len(members & wb_upcs), len(members - wb_upcs)))

    for cid, ups in members_of.items():
        if cid in by_cid or not ups or not ups <= want.keys() or state[cid] not in ("not_reviewed", "decided"):
            continue
        depts = {want[u] for u in ups}
        if len(depts) != 1 or ups & bad_dept:
            continue
        d = depts.pop()
        whole |= ups
        group_rows.append(group_row(cid, d, why[next(iter(ups))][0], len(ups), 0))

    # 4. Every other item
    rest = [u for u in want if u not in whole]
    with engine.connect() as conn:
        live = dm._in_chunks(conn, "SELECT upc, description, department FROM dbo.items WHERE upc IN (SELECT v FROM OPENJSON(:u) WITH (v VARCHAR(400) '$'))", rest)
        bo = dm._in_chunks(conn, "SELECT upc, combo_id, department FROM dbo.dept_mapping_upc_overrides WHERE upc IN (SELECT v FROM OPENJSON(:u) WITH (v VARCHAR(400) '$'))", rest)
        pu = dm._in_chunks(conn, "SELECT upc, department, staged_by FROM dbo.dept_mapping_pending_upc_changes WHERE upc IN (SELECT v FROM OPENJSON(:u) WITH (v VARCHAR(400) '$'))", rest)
        pi = dm._in_chunks(conn, "SELECT upc, department, staged_by FROM dbo.item_master_pending_changes WHERE upc IN (SELECT v FROM OPENJSON(:u) WITH (v VARCHAR(400) '$'))", rest)
    live = live.set_index("upc").to_dict("index") if not live.empty else {}
    bo = bo.set_index("upc").to_dict("index") if not bo.empty else {}
    pu = pu.set_index("upc").to_dict("index") if not pu.empty else {}
    pi = pi.set_index("upc").to_dict("index") if not pi.empty else {}
    item_rows = []
    group_decided = {r["combo_id"] for r in group_rows}
    break_first = {}  # combo_id -> items that need their group broken out first
    for u in rest:
        d = want[u]
        cid = combo_of.get(u)
        c = info.get(cid) if cid is not None else None
        s = state.get(cid)
        it = live.get(u)
        row = {"UPC": u, "Description": (it or {}).get("description"), "Group": c["label"] if c else "(no group)",
               "Area": _where(s, c["tier"], c["n_evidence"]) if c else "Item master", "Now": (it or {}).get("department"),
               "Will be": d, "From (old workbook)": why[u][0], "Goes to": "", "Result": "", "combo_id": cid}
        if it is None:
            row.update(Result=f"{SKIP} — not in the item master")
        elif u in bad_dept:
            row.update(Result=f"{SKIP} — “{d}” isn't one of the app's Departments (add it in Settings first)")
        elif s in ("broken_out", "decided_broken_out") and (u in bo or s == "broken_out"):
            row["Goes to"] = "Broken Out item decision"
            current = (pu.get(u) or {}).get("department") or (None if state.get(cid) != info[cid]["decision_state"]
                                                              else (bo.get(u) or {}).get("department"))
            row["Now"] = current or "(undecided)"
            if current == d:
                row["Result"] = f"{DONE} — already {'staged' if u in pu else 'decided'} as {d}"
            else:
                row["Result"] = STAGE
        elif why[u][1] < 3 and c is not None and cid not in group_decided:
            # The group isn't item-by-item yet: break it out, then decide the item there.
            row["Goes to"] = "Broken Out item decision (group broken out first)"
            row["Now"] = "(undecided)"
            row["Result"] = STAGE
            break_first.setdefault(cid, 0)
            break_first[cid] += 1
        else:
            # Only a real UPC override from the workbook (Final UPC Overrides), an
            # item in no group, or one that differs from its own group's decision.
            row["Goes to"] = "UPC override"
            p = pi.get(u)
            if it["department"] == d and not p:
                row["Result"] = f"{DONE} — already {d}"
            elif p and p["department"] == d:
                row["Result"] = f"{DONE} — already staged ({p['staged_by']})"
            elif p and p["staged_by"] != actor:
                row["Result"] = f"{SKIP} — {p['staged_by']} already has a change staged for this item"
            else:
                row["Result"] = STAGE
        item_rows.append(row)

    for cid, n in break_first.items():
        c = info[cid]
        move_rows.append({"Group (old workbook)": "", "Asked for": f"Break Out — needed for {n:,} item(s) a person decided",
                          "App group": c["label"], "Now": _where(state[cid], c["tier"], c["n_evidence"]), "Result": MOVE,
                          "combo_id": cid, "move": "to_broken_out", "n_upcs_total": int(c["n_upcs_total"] or 0)})
    # 5. Take back: something the downloaded workbook showed as already staged,
    #    that the uploaded copy no longer asks for — only if the person
    #    uploading staged it (anyone else's staged work is never taken back).
    take_rows = []
    wanted_groups = {app_group(g["key"])["combo_id"] for g in extracted["groups"] if app_group(g["key"]) is not None}
    wanted_items = {it["upc"] for it in extracted["items"]}
    staged = extracted.get("staged") or []
    if staged:
        su = [x["upc"] for x in staged if x["upc"]]
        with engine.connect() as conn:
            pu_all = dm._in_chunks(conn, "SELECT upc, combo_id, department, staged_by FROM dbo.dept_mapping_pending_upc_changes WHERE upc IN "
                                   "(SELECT v FROM OPENJSON(:u) WITH (v VARCHAR(400) '$'))", su)
            pi_all = dm._in_chunks(conn, "SELECT upc, department, staged_by, description FROM dbo.item_master_pending_changes WHERE upc IN "
                                   "(SELECT v FROM OPENJSON(:u) WITH (v VARCHAR(400) '$'))", su)
        pu_all = pu_all.set_index("upc").to_dict("index") if not pu_all.empty else {}
        pi_all = pi_all.set_index("upc").to_dict("index") if not pi_all.empty else {}
    for x in staged:
        row = {"What": "", "Group": "", "UPC": x["upc"], "Staged": x["department"], "Staged by": x["staged_by"],
               "Result": "", "combo_id": None}
        if x["kind"] == "group":
            c = app_group(x["key"])
            if c is None:
                continue
            cid = c["combo_id"]
            now = pend_groups.get(cid)
            row.update(What=f"Group decision ({_where(c['decision_state'], c['tier'], c['n_evidence'])})", Group=c["label"], combo_id=cid)
            if cid in wanted_groups or not now or now[0] != x["department"]:
                continue  # still asked for, or no longer staged as it was
            if now[1] != actor:
                row["Result"] = f"{SKIP} — staged by {now[1]}; only your own changes can be taken back"
            else:
                row["Result"] = TAKE
        else:
            u = x["upc"]
            now = (pu_all if x["kind"] == "item" else pi_all).get(u)
            if u in wanted_items or not now or now["department"] != x["department"]:
                continue
            cid = now.get("combo_id")
            row.update(What="Broken Out item decision" if x["kind"] == "item" else "UPC override",
                       Group=info[cid]["label"] if cid in info else "", combo_id=cid)
            row["Result"] = TAKE if now["staged_by"] == actor else f"{SKIP} — staged by {now['staged_by']}; only your own changes can be taken back"
        take_rows.append(row)

    # 6. The workbook was downloaded a while ago? A change to a group or item
    #    that has changed in the app since (someone decided, pushed, moved or
    #    staged it) is skipped — never overwritten. Older workbooks from before
    #    the app have no state sheet and aren't checked.
    snap = extracted.get("state") or {}
    if snap.get("groups") or snap.get("items"):
        key_of = {c["combo_id"]: k for k, c in by_key.items()}
        src_label = {k: lb for lb, k in labels.items()}

        def gkey(cid):
            k = key_of.get(cid)
            return (src_label.get(k[0], k[0].upper()), *k[1:]) if k else None

        def now_group(cid):
            c = info[cid]
            return _group_state(c["decision_state"], c["decided_department"], pend_groups.get(cid))

        with engine.connect() as conn:
            all_upcs = [r["UPC"] for r in item_rows]
            cur_items = dm._in_chunks(conn, "SELECT o.upc, o.department, o.decided_via, p.department AS staged, p.staged_by "
                                            "FROM dbo.dept_mapping_upc_overrides o LEFT JOIN dbo.dept_mapping_pending_upc_changes p "
                                            "ON p.upc = o.upc WHERE o.upc IN (SELECT v FROM OPENJSON(:u) WITH (v VARCHAR(400) '$'))",
                                     all_upcs)
        cur_items = {r["upc"]: _item_state(r["department"], r["decided_via"], (r["staged"], r["staged_by"]) if r["staged"] else None)
                     for r in (cur_items.to_dict("records") if not cur_items.empty else [])}
        n_conflicts = 0

        def conflict(row, then, now):
            nonlocal n_conflicts
            n_conflicts += 1
            row["Result"] = f"{CONFLICT} (then: {_describe_state(then)}; now: {_describe_state(now)}) — download a fresh copy to change it"

        for row in group_rows:
            if row["Result"] == STAGE:
                then = snap["groups"].get(gkey(row["combo_id"]))
                if then is not None and then != now_group(row["combo_id"]):
                    conflict(row, then, now_group(row["combo_id"]))
        for row in move_rows:
            if row["Result"] == MOVE and row.get("combo_id") is not None and row.get("Group (old workbook)"):
                then = snap["groups"].get(gkey(row["combo_id"]))
                if then is not None and then != now_group(row["combo_id"]):
                    conflict(row, then, now_group(row["combo_id"]))
        for row in item_rows:
            if row["Result"] == STAGE:
                then = snap["items"].get(row["UPC"])
                if then is not None and row["UPC"] in cur_items and then != cur_items[row["UPC"]]:
                    conflict(row, then, cur_items[row["UPC"]])
                elif then is None and row.get("combo_id") is not None:
                    # item wasn't in a Broken Out group at download: its group has moved since
                    g_then = snap["groups"].get(gkey(row["combo_id"]))
                    if g_then is not None and g_then.split("|")[0] != info[row["combo_id"]]["decision_state"]:
                        conflict(row, g_then, now_group(row["combo_id"]))

    choices = _choices(extracted, app_group, app_place, info, group_rows, move_rows, item_rows)
    waiting = {c["combo_id"] for c in choices if c["kind"] != "note"}
    for rows_, done_ in ((move_rows, (MOVE,)), (group_rows, (STAGE,)), (item_rows, (STAGE,))):
        for r in rows_:
            if r.get("combo_id") is not None and int(r["combo_id"]) in waiting and r["Result"] in done_:
                r["Result"] = WAIT

    cols_m = ["Group (old workbook)", "Asked for", "App group", "Now", "Result", "combo_id", "move", "n_upcs_total"]
    cols_g = ["Area", "Group", "Items", "Workbook items", "Also covers", "Now", "Will be", "Result", "From (old workbook)", "combo_id", "tier",
              "n_evidence", "source_key", "n_upcs_total"]
    cols_i = ["UPC", "Description", "Area", "Group", "Now", "Will be", "Goes to", "Result", "From (old workbook)", "combo_id"]
    return {"moves": pd.DataFrame(move_rows, columns=cols_m), "groups": pd.DataFrame(group_rows, columns=cols_g),
            "items": pd.DataFrame(item_rows, columns=cols_i),
            "take_back": pd.DataFrame(take_rows, columns=["What", "Group", "UPC", "Staged", "Staged by", "Result", "combo_id"]),
            "choices": pd.DataFrame(choices, columns=CHOICE_COLS),
            "downloaded_at": (extracted.get("state") or {}).get("downloaded_at")}


CHOICE_COLS = ["combo_id", "Group", "kind", "Workbook says", "App has", "Staged now", "Other option", "alt_action",
               "options"]


KEEP_WHAT = "Keep (no change)"


def _row_words(r) -> str:
    return f"“{r['spelled']}” ({r['items']:,} item{'s' if r['items'] != 1 else ''}, {r['place']} sheet): {r['what']}"


def _choices(extracted, app_group, app_place, info, group_rows, move_rows, item_rows) -> list:
    """Where the old workbook and the app don't agree about a group the
    import touches: a list of options — every workbook decision and the
    app's own. Each option: {"who", "text", "goto" (the steps that get
    there), "is" (how to tell it's in effect)}. kind "note" rows aren't
    questions: a group listed more than once where only one row decides
    anything (the import follows that row) — shown on the group's card."""
    out, seen = [], set()
    placements = extracted.get("placements") or {}
    g_by = {r["combo_id"]: r for r in group_rows}
    items_of = {}  # the workbook's Broken Out item decisions, per group: staged only if that option is picked
    for r in item_rows:
        if (r.get("combo_id") is not None and r["Result"] == STAGE and str(r["Goes to"]).startswith("Broken Out item decision")
                and str(r["From (old workbook)"]).startswith(("Old Department UPC Overrides", "Old Decided UPC Overrides",
                                                               "Old Final UPC Overrides"))):
            items_of.setdefault(r["combo_id"], []).append(
                [r["UPC"], r["Will be"], r["Description"], r["Group"], r["From (old workbook)"]])

    def items_opt(cid, text_):
        its = items_of.get(cid) or []
        depts = sorted({i[1] for i in its})
        return {"who": "workbook", "from": "UPC Overrides sheet", "text": text_.format(n=f"{len(its):,}", d=", ".join(depts)),
                "stages": f"Adds {len(its):,} item decisions to Pending Changes ({', '.join(depts)}). The group stays in Broken Out.",
                "goto": "stage_items", "is": {}, "items": its}

    def items_words(cid):
        its = items_of.get(cid) or []
        if not its:
            return ""
        depts = sorted({i[1] for i in its})
        return (f" It also decided {len(its):,} of its items one by one on the UPC Overrides sheet "
                f"({', '.join(depts)}).")

    def listed(key):
        rows = placements.get(key) or []
        if len(rows) < 2:
            return ""
        return (f" The workbook lists this group {len(rows)} times, because the distributor spelled its Department "
                "differently (" + "; ".join(f"“{r['spelled']}”: {r['items']:,} item{'s' if r['items'] != 1 else ''} on the "
                                            f"{r['place']} sheet" for r in rows) + ").")

    def app_long(cid):
        c, pl = info[cid], app_place[cid]
        n = int(c.get("n_upcs_total") or 0)
        return (f"In the app it's one group of {n:,} items, " + {
            "Broken Out": "in Broken Out — being decided item by item.",
            "Decided item by item": "decided item by item.",
            "Decided": f"Decided as {c['decided_department']}."}.get(pl, f"in {pl}, not decided yet."))

    def add(cid, kind, says, has, options):
        if cid in seen:
            return
        seen.add(cid)
        out.append({"combo_id": int(cid), "Group": info[cid]["label"], "kind": kind, "Workbook says": says, "App has": has,
                    "Staged now": "Nothing yet", "Other option": "", "alt_action": "", "options": json.dumps(options)})

    def app_words(cid):
        c, pl = info[cid], app_place[cid]
        if pl == "Decided":
            return f"Decided as {c['decided_department']}"
        return {"Broken Out": "In Broken Out, being decided item by item",
                "Decided item by item": "Decided, item by item"}.get(pl, f"In {pl}, not decided yet")

    nice = {"approved group": "approved the whole group as", "manually overridden group": "Department changed by hand to",
            "manual override": "Department changed by hand to", "Move to Decided Combos": "moved to Decided as"}
    opt = lambda who, text_, goto, is_, from_="", stages="": {"who": who, "from": from_, "text": text_, "stages": stages,
                                                              "goto": goto, "is": is_}

    def queue(cid):
        c = info[cid]
        return "Unmatched" if dm.origin_tier(c["tier"], c["n_evidence"]) == "unmatched" else "Crosswalk"

    def n_of(cid):
        return f"{int(info[cid].get('n_upcs_total') or 0):,}"

    # 1. a group decision made on a sheet that isn't where the app has the group
    for g in extracted["groups"]:
        c = app_group(g["key"])
        if c is None:
            continue
        cid, wb_pl = c["combo_id"], _sheet_of(g["from"])
        if not wb_pl or _same_place(wb_pl, app_place[cid]):
            continue
        row = g_by.get(cid) or {}
        what = g["from"].split(": ", 1)[-1]
        says = f"{wb_pl} sheet: {nice.get(what, what + ' →')} {g['department']}"
        d = g["department"]
        why = (f"In the old workbook this group is on the {wb_pl} sheet, {nice.get(what, what + ' →')} {d}."
                + items_words(cid) + listed(g["key"]))
        if app_place[cid] == "Broken Out":
            options = [opt("workbook", f"Approve the whole group as {d}.", f"send_back_approve:{d}", {"pending": d}, f"{wb_pl} sheet",
                           f"Adds {d} for all {n_of(cid)} items to Pending Changes, and moves the group out of Broken Out "
                           f"back to {queue(cid)}.")]
            if items_of.get(cid):
                options.append(items_opt(cid, "Decide its {n} items one by one, as the workbook did."))
            options.append(opt("app", "Leave it in Broken Out.", "", {"app": True}, "",
                               "Adds nothing. Its items are still decided one by one in Broken Out."))
            add(cid, "place", why, app_long(cid), options)
        elif app_place[cid] == "Decided item by item" and str(row.get("Result", "")).startswith(STAGE):
            add(cid, "place", why, app_long(cid), [
                opt("workbook", f"Decide the whole group as {d}.", f"stage_group:{d}", {"pending": d}, f"{wb_pl} sheet",
                    f"Adds {d} for all {n_of(cid)} items to Pending Changes, replacing its item decisions."),
                opt("app", "Keep its item-by-item decisions.", "", {"app": True}, "", "Adds nothing.")])
    # 2. a Return the workbook asked for, skipped because it has item decisions there
    for m in move_rows:
        cid = m.get("combo_id")
        if cid is None or not m.get("Group (old workbook)"):
            continue
        if str(m["Result"]).startswith(SKIP) and "kept in Broken Out" in str(m["Result"]):
            why = (f"The old workbook's Broken Out sheet asks to {m['Asked for'].split(': ', 1)[-1]} for this group"
                   + (", but it also decides its items one by one." if items_of.get(cid) else ".") + items_words(cid))
            options = [opt("workbook", "Return it for review.", "send_back", {"not_state": "broken_out"}, "Broken Out sheet",
                           f"Moves it back to {queue(cid)} now. Nothing is added to Pending Changes.")]
            if items_of.get(cid):
                options.append(items_opt(cid, "Decide its {n} items one by one, as the workbook did."))
            options.append(opt("app", "Leave it in Broken Out.", "", {"app": True}, "", "Adds nothing."))
            add(cid, "place", why, app_long(cid), options)
    # 3. listed more than once (the distributor spelled its Department two ways)
    for key, rows in placements.items():
        if len(rows) < 2:
            continue
        c = app_group(key)
        if c is None:
            continue
        cid = c["combo_id"]
        decisions = [r for r in rows if r["what"] != KEEP_WHAT]
        mv = next((m for m in move_rows if m.get("combo_id") == cid and m["Result"] == MOVE), None)
        gr = g_by.get(cid)
        touched = mv is not None or (gr is not None and str(gr["Result"]).startswith(STAGE))
        says = ("Listed " + (f"{len(rows)} times" if len(rows) > 2 else "twice") + ", spelled differently — "
                + " · ".join(_row_words(r) for r in rows))
        if len({r["what"] for r in decisions}) <= 1:
            if touched and cid not in seen:  # only one row decides: followed, and noted on the card
                seen.add(cid)
                spelled = " and ".join(f"“{r['spelled']}”" for r in rows)
                short = (f"Listed {'twice' if len(rows) == 2 else f'{len(rows)} times'} in the old workbook (spelled {spelled}). "
                         + (f"Only the {decisions[0]['items']:,}-item row made a change, so that's what was done."
                            if decisions else "No row made a change."))
                out.append({"combo_id": int(cid), "Group": info[cid]["label"], "kind": "note", "Workbook says": short,
                            "App has": app_words(cid), "Staged now": f"Followed the row that decides: {decisions[0]['what']}"
                            if decisions else "Nothing", "Other option": "", "alt_action": "", "options": "[]"})
            continue
        # Two or more rows decide different things: each of them, or the app's.
        options = []
        for r in sorted({r["what"]: r for r in decisions}.values(), key=lambda r: -r["items"]):
            w = r["what"]
            src = f"row “{r['spelled']}”, {r['items']:,} item{'s' if r['items'] != 1 else ''}, {r['place']} sheet"
            if w.startswith(("Change to ", "Approve as ")):
                dept = w.split(" ", 2)[-1]
                options.append(opt("workbook", f"Decide the whole group as {dept}.", f"stage_group:{dept}", {"pending": dept}, src,
                                   f"Adds {dept} for all {n_of(cid)} items to Pending Changes."))
            elif w.startswith("Send to Broken Out"):
                options.append(opt("workbook", "Send it to Broken Out.", "break_out", {"state": "broken_out"}, src,
                                   "Moves it now, to be decided item by item. Nothing is added to Pending Changes."))
            else:
                options.append(opt("workbook", "Return it for review.", "send_back", {"state": "not_reviewed"}, src,
                                   f"Moves it back to {queue(cid)} now. Nothing is added to Pending Changes."))
        options.append(opt("app", f"Leave it as it is: {app_words(cid)}.", "", {"app": True}, "", "Adds nothing."))
        why = (f"The old workbook lists this group {len(rows)} times, because the distributor spelled its Department differently, "
               "and the rows don't agree: " + "; ".join(f"“{r['spelled']}” ({r['items']:,} item{'s' if r['items'] != 1 else ''}, "
                                                        f"{r['place']} sheet) says {r['what']}" for r in rows) + ".")
        add(cid, "twice", why, app_long(cid), options)
    return out


def summary(p: dict) -> pd.DataFrame:
    rows = []
    tb = p.get("take_back", pd.DataFrame(columns=["What", "Result"]))
    for what, df, area_col in (("Moves", p["moves"], "Now"), ("Group decisions", p["groups"], "Area"),
                               ("Item decisions / UPC overrides", p["items"], "Area"), ("Taken back", tb, "What")):
        if df.empty:
            continue
        g = df.assign(Outcome=df["Result"].str.split(" — ").str[0]).groupby([area_col, "Outcome"]).size()
        for (area, outcome), n in g.items():
            rows.append({"What": what, "Where in the app": area or "(not found)", "Outcome": outcome, "Count": int(n)})
    return pd.DataFrame(rows, columns=["What", "Where in the app", "Outcome", "Count"])


def report_excel(p: dict, title: str) -> bytes:
    """Everything in one file: a summary, then every line by area."""
    buf = io.BytesIO()
    hide = ["combo_id", "move", "tier", "n_evidence", "source_key", "n_upcs_total"]  # internal columns
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        pd.DataFrame({"": [title]}).to_excel(xw, sheet_name="Summary", index=False, header=False)
        summary(p).to_excel(xw, sheet_name="Summary", index=False, startrow=2)
        p["moves"].drop(columns=hide, errors="ignore").to_excel(xw, sheet_name="Moves", index=False)
        p["groups"].drop(columns=hide, errors="ignore").to_excel(xw, sheet_name="Group decisions", index=False)
        if len(p.get("take_back", [])):
            p["take_back"].drop(columns=hide, errors="ignore").to_excel(xw, sheet_name="Taken back", index=False)
        items = p["items"].drop(columns=hide, errors="ignore")
        for area in ["Crosswalk", "Unmatched", "Broken Out", "Decided", "Item master"]:
            part = items[items["Area"] == area]
            if not part.empty:
                part.to_excel(xw, sheet_name=f"Items - {area}", index=False)
        for ws in xw.sheets.values():
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(60, max(10, *(len(str(c.value or "")) for c in col[:200])))
    return buf.getvalue()


def apply(engine, p: dict, actor: str, is_admin: bool, file_name: str) -> dict:
    with dm.activity_via("Department workbook upload"):
        return _apply(engine, p, actor, is_admin, file_name)


def _record_choices(engine, p: dict, actor: str, file_name: str) -> int:
    ch = p.get("choices")
    if ch is None or ch.empty:
        return 0
    with engine.begin() as conn:
        conn.execute(text("DELETE FROM dbo.import_choices WHERE status IN ('open', 'note') AND combo_id IN "
                          "(SELECT v FROM OPENJSON(:c) WITH (v INT '$'))").bindparams(bindparam("c", type_=dm.JSON_LIST)),
                     {"c": [int(x) for x in ch["combo_id"]]})
        conn.execute(text(
            "INSERT INTO dbo.import_choices (file_name, combo_id, label, kind, workbook_says, app_has, applied, alternative, "
            "alt_action, created_by, options_json, status) VALUES (:f, :c, :l, :k, :w, :a, :s, :o, :x, :u, :oj, :st)"),
            [{"f": file_name, "c": int(r["combo_id"]), "l": r["Group"], "k": r["kind"], "w": r["Workbook says"],
              "a": r["App has"], "s": r["Staged now"], "o": r["Other option"], "x": r["alt_action"], "u": actor,
              "oj": r["options"], "st": "note" if r["kind"] == "note" else "open"}
             for r in ch.to_dict("records")])
    return len(ch)


def _apply(engine, p: dict, actor: str, is_admin: bool, file_name: str) -> dict:
    """Does everything marked to be done, the same way the app's own buttons
    do it (each group's step is on the top-bar Undo like any other action).
    A safety snapshot is taken first."""
    import uuid
    click = "oldwb-" + uuid.uuid4().hex[:8]

    def note(why: str) -> str:
        """Kept on the staged change, and on the decision once it's pushed."""
        return f"{NOTE_PREFIX} ({file_name}) — {why}"[:400]
    snap = dm.take_snapshot(engine, actor, label=f"Before importing {file_name}", kind="safety_import")
    done = {"snapshot": snap, "moved": 0, "groups": 0, "item_decisions": 0, "suggested": 0, "overrides": 0, "blocked": 0}

    def logged(cid, label, desc, fn):
        before = dm.capture_combo_state(engine, cid)
        fn()
        dm.log_action(engine, actor, cid, label, desc, before, click_id=f"{click}-{cid}")

    for m in p["moves"][p["moves"]["Result"] == MOVE].to_dict("records"):
        cid = int(m["combo_id"])
        snapshot = dm.get_combo_snapshot(engine, cid)
        desc = {"to_broken_out": "Broken Out to UPC-Level", "to_review": "Sent back for review",
                "reopen": "Send Back to Broken Out"}[m["move"]]

        def do(m=m, cid=cid):
            if m["move"] == "to_broken_out":
                dm.break_out_combo(engine, cid, actor, upc_decisions=None)
            elif m["move"] == "reopen":
                dm.reopen_broken_out(engine, cid, actor)
            elif dm.get_combo_snapshot(engine, cid).get("decision_state") in ("broken_out", "decided_broken_out"):
                dm.revert_broken_out_combo(engine, cid, actor)
            else:
                dm.revert_combo(engine, cid, actor)
            dm.clear_pending_for_combo(engine, cid)
        logged(cid, m["App group"], desc, do)
        src, label = m["App group"].split(" — ", 1)
        dm.record_recent_move(engine, cid, src.lower(), label, int(m.get("n_upcs_total") or 0), desc, snapshot, actor,
                              origin_note=note(m["Asked for"]))
        done["moved"] += 1

    for g in p["groups"][p["groups"]["Result"] == STAGE].to_dict("records"):
        cid = int(g["combo_id"])
        label = g["Group"].split(" — ", 1)[1]
        def stage(g=g, cid=cid, label=label):
            dm.upsert_combo_suggestion(engine, cid, dm.origin_tier(g["tier"], g["n_evidence"]), g["Will be"],
                                       g["source_key"], label, int(g["n_upcs_total"]), actor, is_admin=is_admin)
            with engine.begin() as conn:
                conn.execute(text("UPDATE dbo.dept_mapping_pending_changes SET origin_note = :n WHERE combo_id = :c"),
                             {"n": note(g["From (old workbook)"]), "c": cid})
        logged(cid, g["Group"], f"Suggested {g['Will be']} (from {file_name})", stage)
        done["groups"] += 1

    items = p["items"][p["items"]["Result"] == STAGE]
    for cid, grp in items[items["Goes to"].str.startswith("Broken Out item decision")].groupby("combo_id"):
        cid = int(cid)
        glabel = grp["Group"].iloc[0]
        src, label = glabel.split(" — ", 1)
        changes = {r["UPC"]: {"department": r["Will be"], "combo_id": cid, "label": label,
                              "description": r["Description"], "source_key": src.lower(),
                              "origin_note": note(r["From (old workbook)"])} for r in grp.to_dict("records")}
        res = {}
        logged(cid, glabel, f"Set {len(changes)} item(s) from {file_name}",
               lambda: res.update(dm.stage_broken_out_decisions(engine, changes, actor, is_admin=is_admin)))
        done["item_decisions"] += sum(1 for r in res.values() if r["status"] == "decided")
        done["suggested"] += sum(1 for r in res.values() if r["status"] == "suggested")

    tb = p.get("take_back")
    done["taken_back"] = 0
    if tb is not None and len(tb):
        tb = tb[tb["Result"] == TAKE]
        for r in tb[tb["What"].str.startswith("Group decision")].to_dict("records"):
            cid = int(r["combo_id"])

            def take(cid=cid):
                with engine.begin() as conn:
                    others = conn.execute(text("SELECT COUNT(*) FROM dbo.dept_mapping_combo_suggestions WHERE combo_id = :c "
                                               "AND staged_by <> :a"), {"c": cid, "a": actor}).scalar()
                    if others:
                        return
                    conn.execute(text("DELETE FROM dbo.dept_mapping_pending_changes WHERE combo_id = :c"), {"c": cid})
                    conn.execute(text("DELETE FROM dbo.dept_mapping_combo_suggestions WHERE combo_id = :c"), {"c": cid})
                done["taken_back"] += 1
            logged(cid, r["Group"], f"Took back {r['Staged']} (cleared in {file_name})", take)
        items_tb = tb[tb["What"] == "Broken Out item decision"]
        for cid, grp in items_tb.groupby("combo_id"):
            cid = int(cid)
            upcs = grp["UPC"].tolist()

            def take_items(upcs=upcs):
                dm.delete_pending_upc_changes(engine, upcs)
                done["taken_back"] += len(upcs)
            logged(cid, grp["Group"].iloc[0], f"Took back {len(upcs)} item decision(s) (cleared in {file_name})", take_items)
        ov_tb = tb[tb["What"] == "UPC override"]["UPC"].tolist()
        if ov_tb:
            dm.delete_item_master_pending_many(engine, ov_tb)
            done["taken_back"] += len(ov_tb)

    ov = items[items["Goes to"] == "UPC override"]
    if len(ov):
        with engine.connect() as conn:
            cur = dm._in_chunks(conn, "SELECT upc, description, category, subcategory, brand, pack, size, uom, source_key "
                                      "FROM dbo.items WHERE upc IN (SELECT v FROM OPENJSON(:u) WITH (v VARCHAR(400) '$'))", ov["UPC"].tolist()).set_index("upc")
        changes = {}
        for r in ov.to_dict("records"):
            c = cur.loc[r["UPC"]]
            changes[r["UPC"]] = {"change_type": "edit", "department": r["Will be"], "source_key": c["source_key"],
                                 "origin_note": note(r["From (old workbook)"] + f" · group {r['Group']} ({r['Area']})"),
                                 **{f: (None if pd.isna(c[f]) else c[f]) for f in
                                    ("description", "category", "subcategory", "brand", "pack", "size", "uom")}}
        blocked = dm.save_item_master_pending_bulk(engine, changes, actor)
        done["overrides"] = len(changes) - len(blocked)
        done["blocked"] = len(blocked)
    done["choices"] = _record_choices(engine, p, actor, file_name)
    return done


# ---------------------------------------------------------------------------
# The other direction: the app as it stands, as an old-style workbook that can
# be filled in in Excel and brought back in through the import above.
# ---------------------------------------------------------------------------
GROUP_COLS = ["Source", "Old Department", "Category", "Subcategory", "Total UPCs", "New Department", "Confidence",
              "Evidence UPCs", "Runner-Up Department", "Runner-Up Share", "Manual Override Department", "Action"]
ITEM_COLS = ["UPC", "Source", "Description", "Old Department", "Category", "Subcategory", "Mapping Sheet",
             "New Department", "Decided Via", "Manual Override Department"]
ACTIONS = {
    "Department Mapping Crosswalk": ["Not Yet Reviewed", "Approve", "Break Out to UPC-Level"],
    "Department Mapping Unmatched": ["Not Yet Reviewed", "Approve", "Break Out to UPC-Level"],
    "Department Mapping Broken Out": ["Keep (Broken Out)", "Return to Crosswalk/Unmatched"],
    "Department Mapping Decided": ["Keep (No Change)", "Return to Crosswalk", "Return to Unmatched", "Send to Broken Out"],
    "Decided Broken Out Combos": ["Keep (No Change)", "Move to Undecided (Still Broken Out)", "Return to Crosswalk",
                                  "Return to Unmatched", "Move to Decided Combos (Using New Department)"],
}
HOW_TO = {
    "Department Mapping Crosswalk": "Groups with some evidence, waiting on a person. Set Action to Approve to decide the whole group "
                                    "(Manual Override Department if you want something other than New Department), or Break Out to UPC-Level.",
    "Department Mapping Unmatched": "Groups with no evidence at all. Same as Crosswalk: Approve (with a Manual Override Department) or Break Out.",
    "Department Mapping Broken Out": "Groups being decided item by item; their items are on Department UPC Overrides. "
                                     "Action Return to Crosswalk/Unmatched decides the group as a whole instead.",
    "Department Mapping Decided": "Groups already decided as a whole. Type a Manual Override Department to change one, or pick an Action "
                                  "to send it back (Return to Crosswalk/Unmatched) or Send to Broken Out.",
    "Decided Broken Out Combos": "Groups finished item by item (their items are on Decided UPC Overrides). Type a Manual Override Department "
                                 "to decide the whole group instead, or pick an Action to send it back to Broken Out or for review.",
    "Department UPC Overrides": "One row per item of a Broken Out group. Type a Manual Override Department to decide an item "
                                "(rows marked Manually Reviewed or Staged already count as decided).",
    "Decided UPC Overrides": "One row per item of a finished Broken Out group. Type a Manual Override Department to change one.",
    "Final UPC Overrides": "A real UPC override that always wins for that UPC: type a UPC and a New Department on a new row. "
                           "Rows already here are the app's current UPC overrides.",
}


def export_workbook(engine) -> bytes:
    """The app's current state in the old department workbook's layout:
    every group and Broken Out item, with what's staged in Pending Changes
    already filled in, for editing in Excel and importing back."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    q = lambda sql: pd.read_sql(text(sql), engine)
    src = dict(q("SELECT source_key, source_label FROM dbo.sources").itertuples(index=False))
    combos = q("SELECT combo_id, source_key, raw_department, raw_category, raw_subcategory, tier, n_evidence, purity, "
               "decision_state, decided_department, suggested_department, runner_up_department, "
               "runner_up_share, n_upcs_total FROM dbo.dept_mapping_combos")
    pend = dict(q("SELECT combo_id, department FROM dbo.dept_mapping_pending_changes").itertuples(index=False))
    depts = q("SELECT department FROM dbo.dept_mapping_departments ORDER BY department")["department"].tolist()
    items = q("SELECT o.upc, o.combo_id, o.department, o.decided_via, o.suggested_department, i.description, "
              "p.department AS staged FROM dbo.dept_mapping_upc_overrides o LEFT JOIN dbo.items i ON i.upc = o.upc "
              "LEFT JOIN dbo.dept_mapping_pending_upc_changes p ON p.upc = o.upc")
    finals = q("SELECT mo.upc, mo.department, i.description FROM dbo.manual_overrides mo "
               "LEFT JOIN dbo.items i ON i.upc = mo.upc WHERE ISNULL(mo.department, '') <> ''")
    staged_ov = q("SELECT upc, department, description FROM dbo.item_master_pending_changes "
                  "WHERE change_type = 'edit' AND ISNULL(department, '') <> ''")

    def nz(v):
        return "" if v is None or (isinstance(v, float) and pd.isna(v)) else v

    def pct(v):
        return "" if v is None or pd.isna(v) else f"{v:.0%}"

    def queue(r):
        return "Department Mapping Unmatched" if dm.origin_tier(r.tier, r.n_evidence) == "unmatched" else "Department Mapping Crosswalk"

    sheets = {k: [] for k in list(ACTIONS) + ["Department UPC Overrides", "Decided UPC Overrides", "Final UPC Overrides"]}
    where = {}
    for r in combos.itertuples():
        base = {"Source": src.get(r.source_key) or (r.source_key or "").upper(), "Old Department": nz(r.raw_department),
                "Category": nz(r.raw_category), "Subcategory": nz(r.raw_subcategory), "Total UPCs": int(r.n_upcs_total or 0),
                "New Department": nz(r.decided_department) or nz(r.suggested_department), "Confidence": pct(r.purity),
                "Evidence UPCs": int(r.n_evidence or 0), "Runner-Up Department": nz(r.runner_up_department),
                "Runner-Up Share": pct(r.runner_up_share), "Manual Override Department": "", "Action": ""}
        staged = pend.get(r.combo_id)
        if r.decision_state == "broken_out":
            sheet, base["Action"] = "Department Mapping Broken Out", "Keep (Broken Out)"
        elif r.decision_state == "decided_broken_out":
            sheet, base["Action"] = "Decided Broken Out Combos", "Keep (No Change)"
            base["Manual Override Department"] = staged or ""  # a whole-group decision staged on it
        elif r.decision_state == "decided" or nz(r.decided_department):
            sheet, base["Action"] = "Department Mapping Decided", "Keep (No Change)"
            base["Manual Override Department"] = staged or ""
        else:
            sheet = queue(r)
            base["Action"] = "Approve" if staged else "Not Yet Reviewed"
            base["Manual Override Department"] = staged or ""
        sheets[sheet].append(base)
        where[r.combo_id] = (sheet, r)
    for it in items.itertuples():
        sheet_r = where.get(it.combo_id)
        if not sheet_r or sheet_r[0] not in ("Department Mapping Broken Out", "Decided Broken Out Combos"):
            continue
        sheet, r = sheet_r
        target = "Department UPC Overrides" if sheet == "Department Mapping Broken Out" else "Decided UPC Overrides"
        via = nz(it.decided_via)
        sheets[target].append({
            "UPC": it.upc, "Source": src.get(r.source_key) or (r.source_key or "").upper(), "Description": nz(it.description),
            "Old Department": nz(r.raw_department), "Category": nz(r.raw_category), "Subcategory": nz(r.raw_subcategory),
            "Mapping Sheet": queue(r), "New Department": nz(it.department) or nz(it.suggested_department),
            "Decided Via": "Staged (in Pending Changes)" if nz(it.staged) else ("Needs Review" if via in ("", "not_reviewed") else via),
            "Manual Override Department": nz(it.staged)})
    fin = {r.upc: r for r in finals.itertuples()}
    for r in staged_ov.itertuples():
        fin[r.upc] = r
    for u, r in fin.items():
        sheets["Final UPC Overrides"].append({"UPC": u, "New Department": nz(r.department), "Description": nz(r.description)})

    staged_rows = []
    by_id = {r.combo_id: r for r in combos.itertuples()}
    for cid, who_dept in q("SELECT combo_id, department, staged_by FROM dbo.dept_mapping_pending_changes").set_index("combo_id").iterrows():
        r = by_id.get(cid)
        if r is not None:
            staged_rows.append({"Kind": "group", "Source": src.get(r.source_key) or (r.source_key or "").upper(),
                                "Old Department": nz(r.raw_department), "Category": nz(r.raw_category),
                                "Subcategory": nz(r.raw_subcategory), "UPC": "", "Department": who_dept["department"],
                                "Staged by": who_dept["staged_by"]})
    for r in q("SELECT upc, department, staged_by FROM dbo.dept_mapping_pending_upc_changes").itertuples():
        staged_rows.append({"Kind": "item", "UPC": r.upc, "Department": r.department, "Staged by": r.staged_by})
    for r in q("SELECT upc, department, staged_by FROM dbo.item_master_pending_changes WHERE change_type = 'edit' "
               "AND ISNULL(department, '') <> ''").itertuples():
        staged_rows.append({"Kind": "override", "UPC": r.upc, "Department": r.department, "Staged by": r.staged_by})

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        pd.DataFrame({"How to use this workbook": [
            "This is the app as it stands, in the old department workbook's layout. Make changes here, then upload it in the app:",
            "Department Review > Settings > Import decisions from an old department workbook. You get a full report before anything",
            "happens; anything unchanged (or already staged) is skipped as Already done, so only your edits are staged.",
            "",
            *[f"{name}: {how}" for name, how in HOW_TO.items()],
            "",
            "Departments must be one of the app's Departments (see the Departments sheet). Don't rename sheets or columns.",
            "",
            "Rows already filled in are what's staged in Pending Changes right now. Clear one you staged yourself (set the Action",
            "back, or blank its Manual Override Department) and the upload takes that staged change back. Anyone else's staged",
            "changes are never taken back by an upload.",
        ]}).to_excel(xw, sheet_name="Start Here", index=False)
        pd.DataFrame({"Department": depts}).to_excel(xw, sheet_name="Departments", index=False)
        for name, rows in sheets.items():
            cols = ITEM_COLS if name in ("Department UPC Overrides", "Decided UPC Overrides") else (
                ["UPC", "New Department", "Description"] if name == "Final UPC Overrides" else GROUP_COLS)
            df = pd.DataFrame(rows, columns=cols)
            df.to_excel(xw, sheet_name=name, index=False, startrow=1)  # row 1 = how-to, row 2 = headers
            ws = xw.sheets[name]
            ws["A1"] = HOW_TO[name]
            ws["A1"].font = Font(italic=True, color="555555")
            ws["A1"].alignment = Alignment(wrap_text=False)
            ws.freeze_panes = "A3"
            for i, c in enumerate(cols, 1):
                cell = ws.cell(row=2, column=i)
                editable = c in ("Manual Override Department", "Action") or (name == "Final UPC Overrides" and c in ("UPC", "New Department"))
                cell.font = Font(bold=True, color="FFFFFF" if editable else "000000")
                cell.fill = PatternFill("solid", fgColor="2F5597" if editable else "D9D9D9")
                ws.column_dimensions[get_column_letter(i)].width = 30 if c in ("Description", "Category", "Subcategory") else 18
            last = max(len(df) + 2, 3) + (5000 if name == "Final UPC Overrides" else 0)
            dept_list = f"=Departments!$A$2:$A${len(depts) + 1}"
            for c in ("Manual Override Department", "New Department"):
                if c in cols and (c == "Manual Override Department" or name == "Final UPC Overrides"):
                    dv = DataValidation(type="list", formula1=dept_list, allow_blank=True)
                    ws.add_data_validation(dv)
                    col = get_column_letter(cols.index(c) + 1)
                    dv.add(f"{col}3:{col}{last}")
            if name in ACTIONS:
                dv = DataValidation(type="list", formula1='"' + ",".join(ACTIONS[name]) + '"', allow_blank=True)
                ws.add_data_validation(dv)
                col = get_column_letter(cols.index("Action") + 1)
                dv.add(f"{col}3:{col}{last}")
            if "UPC" in cols:
                for row in ws.iter_rows(min_row=3, max_row=last, min_col=1, max_col=1):
                    row[0].number_format = "@"
        xw.sheets["Start Here"].column_dimensions["A"].width = 140
        pd.DataFrame(staged_rows, columns=["Kind", "Source", "Old Department", "Category", "Subcategory", "UPC", "Department",
                                           "Staged by"]).to_excel(xw, sheet_name=STAGED_SHEET, index=False, startrow=1)
        ws = xw.sheets[STAGED_SHEET]
        ws["A1"] = "What was already staged when this workbook was downloaded — used to take back changes you clear. Don't edit."
        ws.sheet_state = "hidden"
        pend_who = {cid: (d, w) for cid, d, w in q("SELECT combo_id, department, staged_by FROM dbo.dept_mapping_pending_changes").itertuples(index=False)}
        upc_pend = {u: (d, w) for u, d, w in q("SELECT upc, department, staged_by FROM dbo.dept_mapping_pending_upc_changes").itertuples(index=False)}
        state_rows = [{"Kind": "downloaded", "State": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")}]
        for r in combos.itertuples():
            state_rows.append({"Kind": "group", "Source": src.get(r.source_key) or (r.source_key or "").upper(),
                               "Old Department": nz(r.raw_department), "Category": nz(r.raw_category), "Subcategory": nz(r.raw_subcategory),
                               "State": _group_state(r.decision_state, nz(r.decided_department) or None, pend_who.get(r.combo_id))})
        for it in items.itertuples():
            state_rows.append({"Kind": "item", "UPC": it.upc,
                               "State": _item_state(nz(it.department) or None, it.decided_via, upc_pend.get(it.upc))})
        pd.DataFrame(state_rows, columns=["Kind", "Source", "Old Department", "Category", "Subcategory", "UPC", "State"]).to_excel(
            xw, sheet_name=STATE_SHEET, index=False, startrow=1)
        ws = xw.sheets[STATE_SHEET]
        ws["A1"] = "The app's state when this workbook was downloaded — used to spot anything changed in the app since. Don't edit."
        ws.sheet_state = "hidden"
    return buf.getvalue()
